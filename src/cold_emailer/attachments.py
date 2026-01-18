"""Resume attachment management and checksum validation."""

import hashlib
from pathlib import Path
from typing import Any

from cold_emailer.utils import get_logger

logger = get_logger(__name__)


def compute_sha256(file_path: Path) -> str:
    """
    Compute SHA256 checksum of a file.

    Args:
        file_path: Path to file

    Returns:
        SHA256 hex digest string

    Raises:
        FileNotFoundError: If file doesn't exist
        IOError: If file cannot be read
    """
    if not file_path.exists():
        raise FileNotFoundError(f"Resume file not found: {file_path}")

    sha256_hash = hashlib.sha256()
    try:
        with open(file_path, "rb") as f:
            for byte_block in iter(lambda: f.read(4096), b""):
                sha256_hash.update(byte_block)
        return sha256_hash.hexdigest()
    except OSError as e:
        raise OSError(f"Failed to read file {file_path}: {e}") from e


def validate_resume_file(
    file_path: Path, expected_sha256: str | None = None
) -> tuple[bool, str | None]:
    """
    Validate resume file exists and optionally verify checksum.

    Args:
        file_path: Path to resume file
        expected_sha256: Expected SHA256 checksum (optional)

    Returns:
        Tuple of (is_valid, error_message)
        - is_valid: True if file exists and checksum matches (if provided)
        - error_message: Error description if validation fails, None if valid
    """
    if not file_path.exists():
        return False, f"Resume file not found: {file_path}"

    if not file_path.is_file():
        return False, f"Path is not a file: {file_path}"

    # Check file size (safety check)
    file_size_mb = file_path.stat().st_size / (1024 * 1024)
    if file_size_mb > 10:  # 10MB limit
        return False, f"Resume file too large: {file_size_mb:.2f}MB (max 10MB)"

    # Verify checksum if provided
    if expected_sha256:
        try:
            actual_sha256 = compute_sha256(file_path)
            if actual_sha256.lower() != expected_sha256.lower():
                return (
                    False,
                    f"Checksum mismatch: expected {expected_sha256}, got {actual_sha256}",
                )
        except Exception as e:
            return False, f"Failed to compute checksum: {e}"

    return True, None


def list_available_resumes(resumes_dir: Path) -> list[dict[str, Any]]:
    """
    List all available resume files in the resumes directory.

    Args:
        resumes_dir: Directory containing resume PDF files

    Returns:
        List of dictionaries with resume info (filename, size, modified_time)
    """
    resumes: list[dict[str, Any]] = []

    if not resumes_dir.exists() or not resumes_dir.is_dir():
        logger.warning("Resumes directory not found", path=str(resumes_dir))
        return resumes

    for file_path in resumes_dir.glob("*.pdf"):
        if file_path.is_file():
            stat = file_path.stat()
            # Get filename without extension
            resume_id = file_path.stem

            resumes.append(
                {
                    "resume_id": resume_id,
                    "filename": file_path.name,
                    "size_bytes": stat.st_size,
                    "size_mb": round(stat.st_size / (1024 * 1024), 2),
                    "modified_time": stat.st_mtime,
                    "path": str(file_path),
                }
            )

    # Sort by modified time (newest first)
    resumes.sort(key=lambda x: x["modified_time"], reverse=True)  # type: ignore[arg-type, return-value]

    logger.info("Listed available resumes", count=len(resumes))
    return resumes


def find_resume_file(
    resume_id: str, resumes_dir: Path
) -> tuple[bool, str | None, dict[str, Any] | None]:
    """
    Find resume file by ID in the resumes directory.

    The resume_id can be:
    - Just the filename (e.g., "RES-001")
    - Filename with extension (e.g., "RES-001.pdf")

    Args:
        resume_id: Resume identifier (filename without or with extension)
        resumes_dir: Directory containing resume PDF files

    Returns:
        Tuple of (is_valid, error_message, resume_info)
        - is_valid: True if resume file found and valid
        - error_message: Error description if invalid, None if valid
        - resume_info: Dictionary with resume info (path, sha256) if found, None otherwise
    """
    if not resumes_dir.exists():
        return False, f"Resumes directory not found: {resumes_dir}", None

    if not resumes_dir.is_dir():
        return False, f"Resumes path is not a directory: {resumes_dir}", None

    # Try different variations of the filename
    possible_names = [
        resume_id,  # Exact match
        f"{resume_id}.pdf",  # With .pdf extension
    ]

    resume_path = None
    for name in possible_names:
        candidate_path = resumes_dir / name
        if candidate_path.exists() and candidate_path.is_file():
            resume_path = candidate_path
            break

    if not resume_path:
        return False, f"Resume file not found: {resume_id} (searched in {resumes_dir})", None

    # Validate the file
    is_valid, error_msg = validate_resume_file(resume_path)
    if not is_valid:
        return False, error_msg or "Resume file validation failed", None

    # Compute checksum
    try:
        sha256 = compute_sha256(resume_path)
        logger.info(
            "Found resume file",
            resume_id=resume_id,
            path=str(resume_path),
            sha256=sha256,
        )
    except Exception as e:
        return False, f"Failed to compute checksum: {e}", None

    resume_info = {
        "resume_id": resume_id,
        "absolute_path": resume_path.resolve(),
        "sha256": sha256,
    }

    return True, None, resume_info


class ResumeManifest:
    """Manages resume manifest and validation."""

    def __init__(self, manifest_path: Path, base_path: Path | None = None) -> None:
        """
        Initialize resume manifest.

        Args:
            manifest_path: Path to resume manifest CSV file
            base_path: Base path for resolving relative paths (defaults to manifest parent)
        """
        self.manifest_path = manifest_path
        self.base_path = base_path or manifest_path.parent
        self.manifest: dict[str, dict[str, Any]] = {}
        self._load_manifest()

    def _load_manifest(self) -> None:
        """Load manifest from CSV or Excel file."""
        import csv

        if not self.manifest_path.exists():
            logger.warning("Resume manifest not found", path=str(self.manifest_path))
            return

        file_ext = self.manifest_path.suffix.lower()

        if file_ext in (".xlsx", ".xls"):
            # Load from Excel
            from openpyxl import load_workbook  # type: ignore[import-untyped]

            workbook = load_workbook(self.manifest_path, read_only=True, data_only=True)
            sheet = workbook.active

            # Read headers
            headers_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
            headers = [str(h).strip() if h else "" for h in headers_row]

            # Create column index map
            col_map = {header: idx for idx, header in enumerate(headers)}

            # Parse rows
            for row_num, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                try:
                    resume_id_col = col_map.get("resume_id")
                    if resume_id_col is None:
                        logger.warning("Missing 'resume_id' column in manifest", row=row_num)
                        continue
                    resume_id = str(row[resume_id_col]).strip() if row[resume_id_col] else ""
                    if not resume_id:
                        continue

                    relative_path_col = col_map.get("relative_path")
                    relative_path = (
                        str(row[relative_path_col]).strip()
                        if relative_path_col is not None and row[relative_path_col]
                        else ""
                    )

                    sha256_col = col_map.get("sha256")
                    sha256 = (
                        str(row[sha256_col]).strip()
                        if sha256_col is not None and row[sha256_col]
                        else None
                    )
                    if sha256:
                        sha256 = sha256 if sha256 else None

                    version_col = col_map.get("version")
                    version = (
                        str(row[version_col]).strip()
                        if version_col is not None and row[version_col]
                        else None
                    )

                    # Resolve absolute path
                    if relative_path:
                        abs_path = (self.base_path / relative_path).resolve()
                    else:
                        abs_path = None

                    # Compute checksum if missing
                    if abs_path and abs_path.exists() and not sha256:
                        try:
                            sha256 = compute_sha256(abs_path)
                            logger.info(
                                "Computed checksum for resume",
                                resume_id=resume_id,
                                sha256=sha256,
                            )
                        except Exception as e:
                            logger.warning(
                                "Failed to compute checksum",
                                resume_id=resume_id,
                                error=str(e),
                            )

                    self.manifest[resume_id] = {
                        "resume_id": resume_id,
                        "relative_path": relative_path,
                        "absolute_path": abs_path,
                        "sha256": sha256,
                        "version": version,
                    }
                except Exception as e:
                    logger.warning("Error parsing manifest row", row=row_num, error=str(e))
                    continue

            workbook.close()
        else:
            # Load from CSV
            try:
                with open(self.manifest_path, encoding="utf-8") as f:
                    reader = csv.DictReader(f)
            except UnicodeDecodeError:
                # Try with different encoding if UTF-8 fails
                with open(self.manifest_path, encoding="utf-8-sig") as f:
                    reader = csv.DictReader(f)
            for row in reader:
                resume_id = row.get("resume_id", "").strip()
                if not resume_id:
                    continue

                relative_path = row.get("relative_path", "").strip()
                sha256 = row.get("sha256", "").strip() or None
                version = row.get("version", "").strip() or None

                # Resolve absolute path
                if relative_path:
                    abs_path = (self.base_path / relative_path).resolve()
                else:
                    abs_path = None

                # Compute checksum if missing
                if abs_path and abs_path.exists() and not sha256:
                    try:
                        sha256 = compute_sha256(abs_path)
                        logger.info(
                            "Computed checksum for resume",
                            resume_id=resume_id,
                            sha256=sha256,
                        )
                    except Exception as e:
                        logger.warning(
                            "Failed to compute checksum",
                            resume_id=resume_id,
                            error=str(e),
                        )

                self.manifest[resume_id] = {
                    "resume_id": resume_id,
                    "relative_path": relative_path,
                    "absolute_path": abs_path,
                    "sha256": sha256,
                    "version": version,
                }

        logger.info("Loaded resume manifest", count=len(self.manifest))

    def get_resume_info(self, resume_id: str) -> dict[str, Any] | None:
        """
        Get resume information by ID.

        Args:
            resume_id: Resume identifier

        Returns:
            Dictionary with resume info or None if not found
        """
        return self.manifest.get(resume_id)

    def validate_resume(self, resume_id: str) -> tuple[bool, str | None, dict[str, Any] | None]:
        """
        Validate resume file exists and checksum matches.

        Args:
            resume_id: Resume identifier

        Returns:
            Tuple of (is_valid, error_message, resume_info)
            - is_valid: True if resume is valid
            - error_message: Error description if invalid, None if valid
            - resume_info: Resume information dictionary if found
        """
        resume_info = self.get_resume_info(resume_id)
        if not resume_info:
            return False, f"Resume ID not found in manifest: {resume_id}", None

        abs_path = resume_info.get("absolute_path")
        if not abs_path:
            return False, f"No path specified for resume: {resume_id}", resume_info

        is_valid, error_msg = validate_resume_file(abs_path, resume_info.get("sha256"))
        return is_valid, error_msg, resume_info

    def update_checksum(self, resume_id: str, sha256: str) -> None:
        """
        Update checksum for a resume in memory (does not persist to file).

        Args:
            resume_id: Resume identifier
            sha256: SHA256 checksum
        """
        if resume_id in self.manifest:
            self.manifest[resume_id]["sha256"] = sha256
